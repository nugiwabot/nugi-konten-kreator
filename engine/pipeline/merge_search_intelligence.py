"""
Script to merge 17 search scraper datasets (Google SERP & ChatGPT search),
deduplicate records, apply cross-encoder Reranker and dense vector Embedding,
and save the unified intelligence dataset into C:\\Users\\Nugi\\Downloads.
"""

import os
import sys
import json
import time
import hashlib
import urllib.parse
import urllib.request
import openpyxl

# Configuration
FILES = [
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-29-55-349.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-57-45-641.json",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-57-41-743.json",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-57-37-941.json",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-57-38-082.json",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-57-41-957.json",
    r"C:\Users\Nugi\Downloads\dataset_chatgpt-search-scraper_2026-09-27_17-24-18-591.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-47-32-747.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-47-20-339.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-47-02-633.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-46-43-494.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-46-37-261.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-46-33-951.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-46-30-447.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-46-30-321.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-46-27-539.xlsx",
    r"C:\Users\Nugi\Downloads\dataset_google-search-results-serp-scraper_2026-09-27_17-43-18-849.xlsx"
]

OUTPUT_FILE = r"C:\Users\Nugi\Downloads\dataset_merged_search_intelligence.json"
CACHE_FILE = r"C:\Users\Nugi\Downloads\.merge_embedding_cache.json"

EMBEDDING_URL = "http://192.168.0.114:1234/v1/embeddings"
RERANKER_URL = "http://192.168.0.114:8080/v1/rerank"

TRACKING_PARAMS = {
    'utm_source', 'utm_medium', 'utm_campaign', 'utm_term', 'utm_content',
    'ref', 'fbclid', 'gclid', 'source', 'platform', 'feature', 'ved'
}

def log(msg):
    print(f"[{time.strftime('%Y-%m-%d %H:%M:%S')}] {msg}", flush=True)

def normalize_url(raw_url):
    if not raw_url:
        return ""
    try:
        raw_url = raw_url.strip()
        parsed = urllib.parse.urlparse(raw_url)
        netloc = parsed.netloc.lower()
        if netloc.startswith("www."):
            netloc = netloc[4:]
        
        # Filter tracking query parameters
        query_params = urllib.parse.parse_qsl(parsed.query)
        filtered_params = [
            (k, v) for k, v in query_params 
            if k.lower() not in TRACKING_PARAMS and not k.lower().startswith('utm_')
        ]
        clean_query = urllib.parse.urlencode(filtered_params)
        path = parsed.path.rstrip('/')
        return urllib.parse.urlunparse((parsed.scheme.lower(), netloc, path, '', clean_query, ''))
    except Exception:
        return raw_url.strip().lower()

def call_reranker(query, documents, batch_size=25, retries=3):
    """
    Call cross-encoder reranker endpoint in batches.
    Returns a dict mapping index in documents -> relevance_score.
    """
    scores = {}
    for start_idx in range(0, len(documents), batch_size):
        chunk = documents[start_idx:start_idx + batch_size]
        payload = {
            "query": query,
            "documents": [doc[:512] for doc in chunk]
        }
        data = json.dumps(payload).encode('utf-8')
        req = urllib.request.Request(RERANKER_URL, data=data, headers={'Content-Type': 'application/json'})
        
        success = False
        for attempt in range(retries):
            try:
                with urllib.request.urlopen(req, timeout=30) as resp:
                    res = json.loads(resp.read().decode('utf-8'))
                    for item in res.get('results', []):
                        rel_idx = item['index']
                        scores[start_idx + rel_idx] = item['relevance_score']
                    success = True
                    break
            except Exception as e:
                log(f"[RERANK] Error on attempt {attempt+1}/{retries}: {e}")
                time.sleep(1.5 * (attempt + 1))
        
        if not success:
            log(f"[RERANK] Warning: Failed to rerank chunk {start_idx} to {start_idx+len(chunk)}")
            for i in range(len(chunk)):
                scores[start_idx + i] = 0.0
                
    return scores

def call_embedding(texts, batch_size=5, retries=3):
    """
    Call embedding endpoint in small batches of 5.
    Returns list of embedding vectors.
    """
    all_embeddings = []
    for i in range(0, len(texts), batch_size):
        chunk = texts[i:i + batch_size]
        payload = {
            "input": [t[:1024] for t in chunk],
            "model": "default"
        }
        data = json.dumps(payload).encode('utf-8')
        req = urllib.request.Request(EMBEDDING_URL, data=data, headers={'Content-Type': 'application/json'})
        
        success = False
        for attempt in range(retries):
            try:
                with urllib.request.urlopen(req, timeout=30) as resp:
                    res = json.loads(resp.read().decode('utf-8'))
                    chunk_embs = [item['embedding'] for item in res.get('data', [])]
                    all_embeddings.extend(chunk_embs)
                    success = True
                    break
            except Exception as e:
                log(f"[EMBED] Error on chunk {i} attempt {attempt+1}/{retries}: {e}")
                time.sleep(2.0 * (attempt + 1))
                
        if not success:
            log(f"[EMBED] Warning: Failed chunk {i} to {i+len(chunk)}, inserting zero vectors")
            for _ in chunk:
                all_embeddings.append([0.0] * 2560)
                
    return all_embeddings

def load_embedding_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_embedding_cache(cache):
    try:
        with open(CACHE_FILE, 'w', encoding='utf-8') as f:
            json.dump(cache, f)
    except Exception as e:
        log(f"[CACHE] Error saving cache: {e}")

def main():
    log("=== STARTING MERGE & INTELLIGENCE PIPELINE ===")
    
    # 1. Parse all files
    raw_records = []
    query_chatgpt_overviews = {}
    
    for fpath in FILES:
        if not os.path.exists(fpath):
            log(f"Warning: File not found: {fpath}")
            continue
            
        fname = os.path.basename(fpath)
        log(f"Parsing file: {fname}...")
        
        if fpath.endswith(".json"):
            with open(fpath, 'r', encoding='utf-8') as fp:
                data = json.load(fp)
                for page in data:
                    q = page.get('search_term', '').strip()
                    for r in page.get('results', []):
                        u = r.get('url', '')
                        if u and u.startswith('http'):
                            raw_records.append({
                                'url': u,
                                'title': (r.get('title') or '').strip(),
                                'snippet': (r.get('description') or '').strip(),
                                'query': q,
                                'position': r.get('position'),
                                'date': r.get('date'),
                                'source_file': fname
                            })
        elif fpath.endswith(".xlsx"):
            wb = openpyxl.load_workbook(fpath, data_only=True)
            sheet = wb.active
            headers = [str(c.value) if c.value is not None else "" for c in sheet[1]]
            
            # Check for ChatGPT Scraper
            if 'query' in headers and any(h.startswith('sources/') for h in headers):
                q_col = headers.index('query')
                text_col = headers.index('text') if 'text' in headers else None
                source_cols = {}
                for i, h in enumerate(headers):
                    if h.startswith('sources/'):
                        parts = h.split('/')
                        s_idx = parts[1]
                        field = parts[2]
                        if s_idx not in source_cols:
                            source_cols[s_idx] = {}
                        source_cols[s_idx][field] = i
                
                for r_idx in range(2, sheet.max_row + 1):
                    q = sheet.cell(row=r_idx, column=q_col+1).value or ""
                    q = str(q).strip()
                    if text_col is not None:
                        synth_text = sheet.cell(row=r_idx, column=text_col+1).value or ""
                        if synth_text and q not in query_chatgpt_overviews:
                            query_chatgpt_overviews[q] = str(synth_text).strip()
                            
                    for s_idx, fmap in source_cols.items():
                        u_col = fmap.get('url')
                        t_col = fmap.get('title')
                        u = sheet.cell(row=r_idx, column=u_col+1).value if u_col is not None else None
                        t = sheet.cell(row=r_idx, column=t_col+1).value if t_col is not None else ""
                        if u and str(u).startswith('http'):
                            raw_records.append({
                                'url': str(u).strip(),
                                'title': str(t).strip() if t else "",
                                'snippet': "",
                                'query': q,
                                'position': int(s_idx) + 1 if s_idx.isdigit() else None,
                                'date': None,
                                'source_file': fname
                            })
            else:
                # Check for Google SERP Scraper
                # Check if it is an error row
                error_cols = [i for i, h in enumerate(headers) if 'error' in h.lower()]
                q_col = None
                for cand in ['search_term', 'keyword', 'query']:
                    if cand in headers:
                        q_col = headers.index(cand)
                        break
                        
                res_cols = {}
                for i, h in enumerate(headers):
                    if h.startswith('results/'):
                        parts = h.split('/')
                        if len(parts) >= 3 and parts[1].isdigit():
                            idx = parts[1]
                            field = parts[2]
                            if idx not in res_cols:
                                res_cols[idx] = {}
                            res_cols[idx][field] = i
                            
                for r_idx in range(2, sheet.max_row + 1):
                    # Check if error cell has value
                    has_err = False
                    for ec in error_cols:
                        if sheet.cell(row=r_idx, column=ec+1).value:
                            has_err = True
                            break
                    if has_err and not res_cols:
                        continue
                        
                    q = sheet.cell(row=r_idx, column=q_col+1).value if q_col is not None else ""
                    q = str(q).strip() if q else ""
                    
                    for r_num, fmap in res_cols.items():
                        u_col = fmap.get('url')
                        t_col = fmap.get('title')
                        d_col = fmap.get('description')
                        p_col = fmap.get('position')
                        dt_col = fmap.get('date')
                        u = sheet.cell(row=r_idx, column=u_col+1).value if u_col is not None else None
                        if u and str(u).startswith('http'):
                            t = sheet.cell(row=r_idx, column=t_col+1).value if t_col is not None else ""
                            d = sheet.cell(row=r_idx, column=d_col+1).value if d_col is not None else ""
                            pos = sheet.cell(row=r_idx, column=p_col+1).value if p_col is not None else None
                            dt = sheet.cell(row=r_idx, column=dt_col+1).value if dt_col is not None else None
                            raw_records.append({
                                'url': str(u).strip(),
                                'title': str(t).strip() if t else "",
                                'snippet': str(d).strip() if d else "",
                                'query': q,
                                'position': pos,
                                'date': str(dt) if dt else None,
                                'source_file': fname
                            })

    log(f"Extracted {len(raw_records)} raw records across {len(FILES)} files.")
    
    # 2. Deduplicate records by normalized URL
    log("Deduplicating records...")
    doc_map = {} # norm_url -> doc_dict
    query_docs_map = {} # query -> list of {norm_url, position}
    
    for rec in raw_records:
        norm_u = normalize_url(rec['url'])
        if not norm_u:
            continue
            
        q = rec['query']
        if q not in query_docs_map:
            query_docs_map[q] = []
        query_docs_map[q].append({
            'norm_url': norm_u,
            'position': rec['position']
        })
        
        if norm_u not in doc_map:
            doc_id = "doc_" + hashlib.md5(norm_u.encode('utf-8')).hexdigest()[:12]
            doc_map[norm_u] = {
                'id': doc_id,
                'url': norm_u,
                'original_urls': set([rec['url']]),
                'title': rec['title'],
                'snippet': rec['snippet'],
                'matched_queries': set([q]) if q else set(),
                'source_files': set([rec['source_file']]),
                'best_position': rec['position'] if isinstance(rec['position'], (int, float)) else 999,
                'dates': set([rec['date']]) if rec['date'] else set()
            }
        else:
            doc = doc_map[norm_u]
            doc['original_urls'].add(rec['url'])
            if q:
                doc['matched_queries'].add(q)
            doc['source_files'].add(rec['source_file'])
            if rec['date']:
                doc['dates'].add(rec['date'])
            if isinstance(rec['position'], (int, float)):
                if doc['best_position'] is None or rec['position'] < doc['best_position']:
                    doc['best_position'] = rec['position']
            # Keep richest title and snippet
            if len(rec['title']) > len(doc['title']):
                doc['title'] = rec['title']
            if len(rec['snippet']) > len(doc['snippet']):
                doc['snippet'] = rec['snippet']

    unique_docs = list(doc_map.values())
    unique_queries = sorted(list(query_docs_map.keys()))
    log(f"Unique normalized documents: {len(unique_docs)}")
    log(f"Unique search queries: {len(unique_queries)}")

    # 3. Apply Cross-Encoder Reranker per Query
    log("=== STAGE 1: APPLYING RERANKER ===")
    query_reranked_results = {}
    
    for q_idx, q in enumerate(unique_queries):
        entries = query_docs_map[q]
        # Deduplicate within query preserving best position
        q_norm_urls = {}
        for entry in entries:
            nu = entry['norm_url']
            pos = entry['position']
            if nu not in q_norm_urls or (isinstance(pos, (int, float)) and pos < q_norm_urls[nu]):
                q_norm_urls[nu] = pos
                
        doc_list = [doc_map[nu] for nu in q_norm_urls.keys()]
        doc_texts = [f"{d['title']} | {d['snippet']}" if d['snippet'] else d['title'] for d in doc_list]
        
        log(f"Reranking Query [{q_idx+1}/{len(unique_queries)}]: '{q}' ({len(doc_list)} candidate documents)...")
        scores = call_reranker(q, doc_texts)
        
        reranked_items = []
        for idx, doc in enumerate(doc_list):
            score = scores.get(idx, 0.0)
            orig_pos = q_norm_urls[doc['url']]
            
            # Record score on master document
            if 'query_scores' not in doc:
                doc['query_scores'] = {}
            doc['query_scores'][q] = round(score, 4)
            
            reranked_items.append({
                'document_id': doc['id'],
                'url': doc['url'],
                'title': doc['title'],
                'snippet': doc['snippet'],
                'relevance_score': round(score, 4),
                'original_position': orig_pos
            })
            
        # Sort descending by relevance score
        reranked_items.sort(key=lambda x: x['relevance_score'], reverse=True)
        for rank_num, item in enumerate(reranked_items, 1):
            item['rank'] = rank_num
            
        query_reranked_results[q] = reranked_items
        log(f"  -> Top 1 result: '{reranked_items[0]['title'][:50]}...' (Score: {reranked_items[0]['relevance_score']})")

    # Update master documents with best matching query and max score
    for doc in unique_docs:
        q_scores = doc.get('query_scores', {})
        if q_scores:
            best_q = max(q_scores, key=q_scores.get)
            doc['best_matching_query'] = best_q
            doc['max_relevance_score'] = q_scores[best_q]
        else:
            doc['best_matching_query'] = None
            doc['max_relevance_score'] = None

    # 4. Apply Dense Vector Embedding
    log("=== STAGE 2: APPLYING EMBEDDINGS ===")
    cache = load_embedding_cache()
    log(f"Existing embeddings in cache: {len(cache)}")
    
    # 4a. Embed all Queries
    log("Embedding 36 Queries...")
    query_embeddings = {}
    queries_to_embed = [q for q in unique_queries if f"query::{q}" not in cache]
    if queries_to_embed:
        embs = call_embedding(queries_to_embed, batch_size=4)
        for q, emb in zip(queries_to_embed, embs):
            cache[f"query::{q}"] = emb
        save_embedding_cache(cache)
        
    for q in unique_queries:
        query_embeddings[q] = cache.get(f"query::{q}", [])
    log("Completed embedding queries.")

    # 4b. Embed all 1,298 Unique Documents
    docs_to_embed = [d for d in unique_docs if f"doc::{d['id']}" not in cache]
    log(f"Documents to embed: {len(docs_to_embed)} (already cached: {len(unique_docs) - len(docs_to_embed)})")
    
    BATCH_SIZE = 5
    t_start = time.time()
    
    for i in range(0, len(docs_to_embed), BATCH_SIZE):
        batch = docs_to_embed[i:i + BATCH_SIZE]
        texts = []
        for d in batch:
            text = f"{d['title']}. {d['snippet']}" if d['snippet'] else d['title']
            texts.append(text)
            
        t_batch_start = time.time()
        embs = call_embedding(texts, batch_size=BATCH_SIZE)
        t_batch = time.time() - t_batch_start
        
        for d, emb in zip(batch, embs):
            cache[f"doc::{d['id']}"] = emb
            
        # Periodic cache save and progress log
        done = i + len(batch)
        elapsed = time.time() - t_start
        rate = done / elapsed if elapsed > 0 else 0.1
        remaining_sec = (len(docs_to_embed) - done) / rate if rate > 0 else 0
        rem_min = int(remaining_sec // 60)
        rem_sec = int(remaining_sec % 60)
        
        if done % 25 == 0 or done == len(docs_to_embed):
            save_embedding_cache(cache)
            log(f"[EMBED PROGRESS] {done}/{len(docs_to_embed)} ({done/len(docs_to_embed)*100:.1f}%) "
                f"- Batch: {t_batch:.1f}s - Avg: {1/rate:.1f}s/doc - ETA: {rem_min}m {rem_sec}s")

    # Assign embeddings to master documents
    for doc in unique_docs:
        doc['embedding'] = cache.get(f"doc::{doc['id']}", [])

    # 5. Format and Save Final Output
    log("=== STAGE 3: FORMATTING FINAL DATASET ===")
    
    formatted_queries = []
    for q_idx, q in enumerate(unique_queries, 1):
        slug = "-".join(q.lower().split())
        formatted_queries.append({
            "query_id": f"q_{q_idx:02d}",
            "query": q,
            "slug": slug,
            "total_results": len(query_reranked_results.get(q, [])),
            "chatgpt_overview": query_chatgpt_overviews.get(q, None),
            "embedding": query_embeddings.get(q, []),
            "reranked_results": query_reranked_results.get(q, [])
        })

    # Prepare serialized documents
    formatted_documents = []
    for doc in unique_docs:
        formatted_documents.append({
            "id": doc['id'],
            "url": doc['url'],
            "original_urls": sorted(list(doc['original_urls'])),
            "title": doc['title'],
            "snippet": doc['snippet'],
            "matched_queries": sorted(list(doc['matched_queries'])),
            "source_files": sorted(list(doc['source_files'])),
            "best_position": doc['best_position'] if doc['best_position'] != 999 else None,
            "dates": sorted(list(doc['dates'])),
            "best_matching_query": doc['best_matching_query'],
            "max_relevance_score": doc['max_relevance_score'],
            "query_scores": doc.get('query_scores', {}),
            "embedding": doc.get('embedding', [])
        })

    final_payload = {
        "metadata": {
            "title": "Nugi Konten Kreator - Merged Search Intelligence Dataset",
            "created_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
            "total_source_files": len(FILES),
            "total_raw_scraped_records": len(raw_records),
            "total_unique_documents": len(formatted_documents),
            "total_unique_queries": len(formatted_queries),
            "total_embedded_queries": len(formatted_queries),
            "total_embedded_documents": sum(1 for d in formatted_documents if d.get('embedding')),
            "embedding_model": "Qwen3-Embedding-4B-Q4_K_M",
            "embedding_endpoint": EMBEDDING_URL,
            "embedding_dimension": 2560,
            "reranker_model": "bge-reranker-v2-m3-q8_0",
            "reranker_endpoint": RERANKER_URL,
            "output_file": OUTPUT_FILE
        },
        "queries": formatted_queries,
        "documents": formatted_documents
    }

    log(f"Writing unified JSON dataset to {OUTPUT_FILE}...")
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as fp:
        json.dump(final_payload, fp, ensure_ascii=False, indent=2)

    file_size_mb = os.path.getsize(OUTPUT_FILE) / (1024 * 1024)
    log(f"SUCCESS! Output saved to: {OUTPUT_FILE}")
    log(f"File size: {file_size_mb:.2f} MB")
    log("=== PROCESSING COMPLETED ===")

if __name__ == "__main__":
    main()
